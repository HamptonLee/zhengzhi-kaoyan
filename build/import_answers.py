#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""导入答案与讲解。

用法:
  python build/import_answers.py --template           # 生成空模板 import/答案导入模板.csv
  python build/import_answers.py --file 答案.csv       # 从 CSV/XLSX 导入
  python build/import_answers.py --file 解析.docx      # 从 DOCX/TXT/PDF 导入（需含「年份-题型（题号）」锚点）
  ... --force                                         # 覆盖已有答案
  ... --no-build                                      # 只写 answers.json，不重建 questions.js

答案文件里的「答案」支持 A / ABD / A、B、D / A,B,D 等写法。
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
Q_JSON = BASE / "data" / "questions.json"
A_JSON = BASE / "data" / "answers.json"
TEMPLATE = BASE / "import" / "答案导入模板.csv"

ANCHOR = re.compile(r"(\d{4})\s*[-－—]?\s*(单选|多选)\s*[.．]?\s*[（(]\s*(\d+)\s*[）)]")
ANS_RE = re.compile(r"(?:【\s*答案\s*】|答案)\s*[:：]?\s*([A-E][A-Ea-e、，,．.\s]{0,10})")
EXPLAIN_SPLIT = re.compile(r"(?:【\s*解析\s*】|【\s*分析\s*】|解析|分析|考点|【\s*考点\s*】)")

HEADER_ALIASES = {
    "id": ["题目id", "id", "题目编号", "编号"],
    "subject": ["学科", "科目"],
    "part": ["部分", "模块"],
    "type": ["题型"],
    "year": ["年份", "年度"],
    "no": ["卷面题号", "题号", "序号"],
    "answer": ["答案", "正确答案"],
    "explain": ["讲解", "解析", "说明", "答案解析"],
}


def norm_answer(raw: str) -> str:
    letters = [c.upper() for c in str(raw or "") if c.upper() in "ABCDE"]
    seen = []
    for c in letters:
        if c not in seen:
            seen.append(c)
    return "".join(sorted(seen))


def read_table(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    rows: list[dict] = []
    if suffix in {".csv", ".txt"} and suffix == ".csv":
        text = None
        for enc in ("utf-8-sig", "gbk", "utf-8"):
            try:
                text = path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise SystemExit(f"无法解码 CSV: {path}")
        reader = csv.DictReader(text.splitlines())
        for r in reader:
            rows.append({(k or "").strip(): (v or "").strip() for k, v in r.items()})
        return rows

    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return rows
        header = [str(v).strip() if v is not None else "" for v in values[0]]
        for line in values[1:]:
            rows.append({header[i]: (str(line[i]).strip() if i < len(line) and line[i] is not None else "")
                         for i in range(len(header))})
        return rows

    raise SystemExit(f"表格导入不支持该格式: {suffix}（请用 CSV 或 XLSX）")


def read_text_blocks(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise SystemExit("读取 PDF 需要 pypdf：python -m pip install pypdf")
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix in {".txt", ".md"}:
        for enc in ("utf-8-sig", "gbk", "utf-8"):
            try:
                return path.read_text(encoding=enc)
            except UnicodeDecodeError:
                continue
        raise SystemExit(f"无法解码文本文件: {path}")
    raise SystemExit(f"不支持的文件格式: {suffix}")


def match_columns(row: dict) -> dict:
    out = {}
    lower = {k.lower(): v for k, v in row.items()}
    for field, aliases in HEADER_ALIASES.items():
        for a in aliases:
            if a in row:
                out[field] = row[a]
                break
            if a.lower() in lower:
                out[field] = lower[a.lower()]
                break
    return out


def build_index(questions: list[dict]) -> tuple[dict, dict]:
    by_id = {q["id"]: q for q in questions}
    by_key: dict[tuple, list[dict]] = {}
    for q in questions:
        by_key.setdefault((q["year"], q["type"], q["no"]), []).append(q)
    return by_id, by_key


def resolve(by_id: dict, by_key: dict, rec: dict) -> tuple[dict | None, str]:
    qid = (rec.get("id") or "").strip()
    if qid and qid in by_id:
        return by_id[qid], "id"
    try:
        key = (int(rec["year"]), rec["type"].strip(), int(rec["no"]))
    except (KeyError, TypeError, ValueError):
        return None, "缺少年份/题型/题号"
    candidates = by_key.get(key) or []
    if len(candidates) == 1:
        return candidates[0], "year+type+no"
    if len(candidates) > 1:
        return None, f"命中 {len(candidates)} 道同号题（卷面题号重复），请在表格里填题目ID"
    return None, "题库中找不到"


def import_table(path: Path, questions: list[dict]) -> tuple[dict, list[str], list[str]]:
    rows = read_table(path)
    by_id, by_key = build_index(questions)
    answers: dict[str, dict] = {}
    missing: list[str] = []
    ok = 0
    for i, row in enumerate(rows, start=2):
        rec = match_columns(row)
        if not (rec.get("answer") or rec.get("explain")):
            continue
        q, how = resolve(by_id, by_key, rec)
        if not q:
            missing.append(f"第{i}行: {how} | {str(row)[:80]}")
            continue
        answers[q["id"]] = {
            "answer": norm_answer(rec.get("answer")),
            "explain": (rec.get("explain") or "").strip() or None,
            "source": "imported",
            "updatedAt": datetime.now().isoformat(timespec="seconds"),
        }
        ok += 1
    note = [f"表格导入：识别 {ok} 条"]
    return answers, missing, note


def import_text(path: Path, questions: list[dict]) -> tuple[dict, list[str], list[str]]:
    text = read_text_blocks(path)
    by_id, by_key = build_index(questions)
    anchors = list(ANCHOR.finditer(text))
    answers: dict[str, dict] = {}
    missing: list[str] = []
    for i, m in enumerate(anchors):
        year, typ, no = int(m.group(1)), m.group(2), int(m.group(3))
        end = anchors[i + 1].start() if i + 1 < len(anchors) else len(text)
        block = text[m.end():end]
        am = ANS_RE.search(block)
        if not am:
            continue
        ans = norm_answer(am.group(1))
        if not ans:
            continue
        rest = block[am.end():].strip()
        em = EXPLAIN_SPLIT.match(rest)
        if em:
            rest = rest[em.end():].strip(" ：:\n\t")
        candidates = by_key.get((year, typ, no)) or []
        if len(candidates) != 1:
            missing.append(f"{year}-{typ}（{no}）: 命中 {len(candidates)} 道（题号重复或有歧义）")
            continue
        q = candidates[0]
        answers[q["id"]] = {
            "answer": ans,
            "explain": rest.strip() or None,
            "source": "imported",
            "updatedAt": datetime.now().isoformat(timespec="seconds"),
        }
    return answers, missing, [f"文本导入：解析出 {len(answers)} 条"]


def write_template(questions: list[dict]) -> None:
    TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
    with TEMPLATE.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["题目ID", "学科", "部分", "题型", "年份", "卷面题号", "题干摘要", "答案", "讲解"])
        for q in questions:
            w.writerow([q["id"], q["subject"], q["part"] or "", q["type"], q["year"], q["no"],
                        q["stem"][:40], "", ""])
    print(f"已生成模板: {TEMPLATE}（{len(questions)} 行）")


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    ap = argparse.ArgumentParser()
    ap.add_argument("--file", type=str, help="答案文件路径（csv/xlsx/docx/txt/pdf）")
    ap.add_argument("--template", action="store_true", help="生成空白答案模板")
    ap.add_argument("--force", action="store_true", help="覆盖已有答案")
    ap.add_argument("--no-build", action="store_true", help="不重建 questions.js")
    args = ap.parse_args()

    payload = json.loads(Q_JSON.read_text(encoding="utf-8-sig"))
    questions = payload["questions"]

    if args.template:
        write_template(questions)
        return 0

    if not args.file:
        ap.error("需要 --file 或 --template")

    path = Path(str(args.file).strip().strip('"').strip("'"))
    if not path.exists():
        raise SystemExit(f"文件不存在: {path}")

    if path.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
        incoming, missing, notes = import_table(path, questions)
    else:
        incoming, missing, notes = import_text(path, questions)

    existing = {}
    if A_JSON.exists():
        existing = json.loads(A_JSON.read_text(encoding="utf-8-sig")).get("answers", {})

    added = updated = skipped = 0
    for qid, val in incoming.items():
        if qid in existing and existing[qid].get("source") == "manual" and not args.force:
            skipped += 1
            continue
        if qid in existing:
            updated += 1
        else:
            added += 1
        existing[qid] = val

    A_JSON.write_text(
        json.dumps({"updatedAt": datetime.now().isoformat(timespec="seconds"), "answers": existing},
                   ensure_ascii=False, indent=1),
        encoding="utf-8",
    )

    for n in notes:
        print(n)
    print(f"新增 {added} 条，覆盖 {updated} 条，跳过 {skipped} 条（手填保护）")
    if missing:
        print(f"\n未匹配 {len(missing)} 条（前 10 条）:")
        for m in missing[:10]:
            print("  -", m)
        if len(missing) > 10:
            print(f"  ... 其余 {len(missing) - 10} 条")

    if not args.no_build:
        sys.path.insert(0, str(BASE / "build"))
        import build_data  # noqa: E402

        return build_data.main()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
